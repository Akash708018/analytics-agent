"""
Unit tests for ingest/merges.py.

Every fixture here is built in a tmp_path so the tests are self-contained and
do not depend on tests/fixtures/ being present or unchanged.
"""

from __future__ import annotations

import re
import zipfile

import pytest
from openpyxl import Workbook

from analytics_agent.ingest.merges import (
    MergeParseError,
    fill_bounded,
    merged_ranges,
    merges_on_rows,
    sheet_names,
)


# --------------------------------------------------------------------------
# fixtures
# --------------------------------------------------------------------------

@pytest.fixture
def spanning_header(tmp_path):
    """
    Row 4: A4='Order', C4='Q3 2024' merged across C4:D4. E4 and F4 are blank
    and are under no merge at all.
    Row 5: the leaf header.
    This is the exact shape that exposes unbounded fill.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Q3 Data"
    ws["A4"] = "Order"
    ws["C4"] = "Q3 2024"
    for col, val in zip("ABCDEF", ["ID", "Region", "Units", "Amt", "Units", "Channel"]):
        ws[f"{col}5"] = val
    ws.merge_cells("C4:D4")
    ws["A6"] = "A1"
    ws["F6"] = "web"
    path = tmp_path / "spanning.xlsx"
    wb.save(path)
    return str(path)


@pytest.fixture
def reordered_tabs(tmp_path):
    """
    Two sheets, then <sheets> order reversed in workbook.xml while the
    sheetN.xml filenames stay put -- what Excel writes when a tab is dragged.
    'Data' keeps C4:D4; 'Notes' keeps A1:B1.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["C4"] = "Q3 2024"
    ws.merge_cells("C4:D4")
    notes = wb.create_sheet("Notes")
    notes["A1"] = "nothing here"
    notes.merge_cells("A1:B1")
    src = tmp_path / "src.xlsx"
    wb.save(src)

    dst = tmp_path / "reordered.xlsx"
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                text = data.decode()
                block = re.search(r"<sheets>(.*?)</sheets>", text, re.S).group(1)
                elems = re.findall(r"<sheet .*?/>", block)
                text = text.replace(block, "".join(reversed(elems)))
                data = text.encode()
            zout.writestr(item, data)
    return str(dst)


# --------------------------------------------------------------------------
# fill_bounded -- FMR F14
# --------------------------------------------------------------------------

def _join(top, bottom):
    return " ".join(v for v in (top, bottom) if v not in (None, ""))


def test_bounded_fill_does_not_leak_past_the_merge(spanning_header):
    """The regression test named in Phase 3 step 4: Channel stays Channel."""
    refs = merged_ranges(spanning_header, "Q3 Data")
    row4 = ["Order", None, "Q3 2024", None, None, None]
    row5 = ["ID", "Region", "Units", "Amt", "Units", "Channel"]

    filled = fill_bounded(row4, refs, row_num_1idx=4)
    names = [_join(t, b) for t, b in zip(filled, row5)]

    assert names == [
        "Order ID",
        "Region",
        "Q3 2024 Units",
        "Q3 2024 Amt",
        "Units",
        "Channel",
    ]
    assert names[5] == "Channel", "Q3 2024 leaked onto a column under no merge"
    assert names[1] == "Region", "Order leaked out of A4 into column B"


def test_bounded_fill_ignores_merges_on_other_rows():
    row = ["a", None, "b", None]
    assert fill_bounded(row, ["B9:D9"], row_num_1idx=4) == row


def test_bounded_fill_pads_rows_shorter_than_the_merge():
    """A row shorter than the merge extent must pad, not raise IndexError."""
    out = fill_bounded(["Order", None, "Q3 2024"], ["C4:F4"], row_num_1idx=4)
    assert out == ["Order", None, "Q3 2024", "Q3 2024", "Q3 2024", "Q3 2024"]


def test_bounded_fill_propagates_an_empty_anchor():
    """An empty anchor stays empty. It is not an invitation to guess."""
    assert fill_bounded(["x", None, None], ["B4:C4"], row_num_1idx=4) == ["x", None, None]


def test_bounded_fill_does_not_mutate_its_input():
    row = ["Order", None, "Q3 2024", None]
    fill_bounded(row, ["C4:D4"], row_num_1idx=4)
    assert row == ["Order", None, "Q3 2024", None]


# --------------------------------------------------------------------------
# merged_ranges -- FMR F11
# --------------------------------------------------------------------------

def test_merged_ranges_matches_normal_mode_ground_truth(spanning_header):
    from openpyxl import load_workbook

    truth = sorted(
        str(r) for r in load_workbook(spanning_header)["Q3 Data"].merged_cells.ranges
    )
    assert merged_ranges(spanning_header, "Q3 Data") == truth


def test_read_only_worksheet_still_has_no_merged_cells(spanning_header):
    """
    Pins the reason this module exists. If a future openpyxl adds the
    attribute, this test fails and the XML path can be reconsidered --
    deliberately, not by accident.
    """
    from openpyxl import load_workbook

    ws = load_workbook(spanning_header, read_only=True)["Q3 Data"]
    with pytest.raises(AttributeError):
        _ = ws.merged_cells


def test_sheet_is_resolved_by_relationship_not_by_position(reordered_tabs):
    """
    Tab order != file numbering. Resolving by position in <sheets> returns
    the wrong sheet's merges here, with no error.
    """
    assert sheet_names(reordered_tabs) == ["Notes", "Data"]
    assert merged_ranges(reordered_tabs, "Data") == ["C4:D4"]
    assert merged_ranges(reordered_tabs, "Notes") == ["A1:B1"]


def test_unknown_sheet_name_lists_what_is_available(spanning_header):
    with pytest.raises(KeyError) as exc:
        merged_ranges(spanning_header, "Sheet1")
    assert "Q3 Data" in str(exc.value)


def test_non_ooxml_file_gives_an_actionable_error(tmp_path):
    fake = tmp_path / "old.xls"
    fake.write_bytes(b"\xd0\xcf\x11\xe0not a zip")
    with pytest.raises(MergeParseError) as exc:
        merged_ranges(str(fake), "Sheet1")
    assert "re-save" in str(exc.value).lower()


def test_sheet_with_no_merges_returns_empty(tmp_path):
    wb = Workbook()
    wb.active.title = "Plain"
    wb.active["A1"] = "x"
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    assert merged_ranges(str(path), "Plain") == []


def test_merges_on_rows_filters_by_row_span():
    refs = ["C4:D4", "A1:B1", "B7:C9"]
    assert merges_on_rows(refs, [4]) == ["C4:D4"]
    assert merges_on_rows(refs, [8]) == ["B7:C9"]
    assert merges_on_rows(refs, [1, 4]) == ["A1:B1", "C4:D4"]
    assert merges_on_rows(refs, [5]) == []
